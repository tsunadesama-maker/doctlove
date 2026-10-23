"""
DocuTrad AI - Service de Traitement et Traduction DOCX
Lecture fine des paragraphes Word, préservation des styles (runs, gras, italique, souligné),
traduction des tableaux en conservant leur structure, ainsi que des en-têtes et pieds de page.
"""
from pathlib import Path
from typing import Dict, Any, Optional, List
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .translation_service import TranslationService
from .file_validation import FileValidationService, ValidationError
from ..config import FREE_TIER_MAX_PAGES

class DOCXService:
    def __init__(self, translation_service: Optional[TranslationService] = None):
        self.translator = translation_service or TranslationService()

    def inspect_docx(self, file_path: Path) -> Dict[str, Any]:
        """Analyse les métadonnées et estime les pages du document Word."""
        page_count = FileValidationService.count_docx_pages(file_path)
        doc = Document(file_path)

        paragraphs_count = len(doc.paragraphs)
        tables_count = len(doc.tables)

        # Échantillon de texte pour détection de langue
        sample_texts = [p.text for p in doc.paragraphs[:10] if p.text.strip()]
        sample_str = " ".join(sample_texts)
        detected_lang = self.translator.detect_language(sample_str)

        return {
            "page_count": page_count,
            "paragraphs_count": paragraphs_count,
            "tables_count": tables_count,
            "detected_lang": detected_lang,
            "is_scanned": False,
            "requires_ocr": False
        }

    def _translate_paragraph_runs(
        self,
        paragraph,
        source_lang: str,
        target_lang: str
    ) -> int:
        """
        Traduit un paragraphe en préservant scrupuleusement la structure des runs
        (gras, italique, souligné, couleur, taille).
        """
        full_text = paragraph.text
        if not full_text.strip():
            return 0

        runs = paragraph.runs
        words_translated = 0

        if not runs:
            # Si pas de runs explicites, traduire directement le texte du paragraphe
            translated = self.translator.translate_chunk(full_text, source_lang, target_lang)
            paragraph.text = translated
            return len(translated.split())

        if len(runs) == 1:
            # 1 seul run : traduction directe en conservant l'objet run et son formatage intact
            run = runs[0]
            if run.text.strip():
                orig_text = run.text
                translated = self.translator.translate_chunk(orig_text, source_lang, target_lang)
                run.text = translated
                words_translated += len(translated.split())
        else:
            # Plusieurs runs : chaque run conserve ses styles (gras, italique, etc.)
            for run in runs:
                if run.text.strip():
                    orig_text = run.text
                    translated = self.translator.translate_chunk(orig_text, source_lang, target_lang)
                    run.text = translated
                    words_translated += len(translated.split())

        return words_translated

    def process_and_translate_docx(
        self,
        input_path: Path,
        output_path: Path,
        source_lang: str,
        target_lang: str,
        is_premium: bool = False
    ) -> Dict[str, Any]:
        """
        Ouvre le document DOCX, traduit les paragraphes, tableaux, en-têtes et pieds de page,
        puis enregistre le document traduit sans altérer les styles Word.
        """
        inspection = self.inspect_docx(input_path)
        page_count = inspection["page_count"]

        # 1. Vérification du plafond des 19 pages
        is_allowed, limit_msg = FileValidationService.check_page_limits(page_count, is_premium)
        if not is_allowed:
            raise ValidationError(limit_msg)

        effective_source = source_lang
        if source_lang == "auto":
            effective_source = inspection["detected_lang"]

        doc = Document(input_path)
        total_words_translated = 0
        translated_paragraphs_count = 0

        # 2. Traduction des paragraphes principaux
        for p in doc.paragraphs:
            words = self._translate_paragraph_runs(p, effective_source, target_lang)
            if words > 0:
                translated_paragraphs_count += 1
                total_words_translated += words

        # 3. Traduction des tableaux (en préservant la structure et les cellules)
        translated_tables_count = 0
        for table in doc.tables:
            translated_tables_count += 1
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        words = self._translate_paragraph_runs(p, effective_source, target_lang)
                        total_words_translated += words

        # 4. Traduction des en-têtes et pieds de page de chaque section
        for section in doc.sections:
            if section.header:
                for p in section.header.paragraphs:
                    words = self._translate_paragraph_runs(p, effective_source, target_lang)
                    total_words_translated += words
            if section.footer:
                for p in section.footer.paragraphs:
                    words = self._translate_paragraph_runs(p, effective_source, target_lang)
                    total_words_translated += words

        # 5. Sauvegarde du nouveau fichier DOCX
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)

        # 6. Vérification que le fichier de sortie est non vide
        if not output_path.exists() or output_path.stat().st_size == 0:
            raise RuntimeError("La génération du document DOCX traduit a échoué (fichier vide).")

        return {
            "success": True,
            "page_count": page_count,
            "translated_paragraphs": translated_paragraphs_count,
            "translated_tables": translated_tables_count,
            "total_words_translated": total_words_translated,
            "source_lang": effective_source,
            "target_lang": target_lang,
            "output_size_bytes": output_path.stat().st_size,
            "output_filename": output_path.name
        }

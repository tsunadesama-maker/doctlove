"""
DocuTrad AI - Service de Validation de Fichiers
Vérification des types réels (magic bytes), limites de taille,
sécurité des chemins, et comptage serveur des pages.
"""
import os
import uuid
from pathlib import Path
from typing import Tuple, Optional
import pymupdf as fitz
from docx import Document

from ..config import (
    ALLOWED_EXTENSIONS,
    MAX_FILE_SIZE_BYTES,
    FREE_TIER_MAX_PAGES,
    UPLOAD_DIR,
    OUTPUT_DIR,
)

class ValidationError(Exception):
    """Exception levée en cas d'échec de validation d'un document."""
    pass

class FileValidationService:
    @staticmethod
    def sanitize_filename(original_name: str) -> str:
        """Nettoie le nom de fichier et génère un identifiant unique sécurisé."""
        safe_name = os.path.basename(original_name).replace("..", "").replace("/", "").replace("\\", "")
        clean_ext = Path(safe_name).suffix.lower()
        if clean_ext not in ALLOWED_EXTENSIONS:
            clean_ext = ".bin"
        unique_id = str(uuid.uuid4())
        return f"{unique_id}{clean_ext}"

    @staticmethod
    def validate_file_bytes(file_bytes: bytes, filename: str) -> Tuple[bool, str, str]:
        """
        Vérifie l'extension, les magic bytes et la taille du fichier.
        Retourne (is_valid, format_detected, error_message).
        """
        # 1. Vérification de la taille
        if len(file_bytes) == 0:
            return False, "", "Le fichier envoyé est vide (0 octet)."
        
        if len(file_bytes) > MAX_FILE_SIZE_BYTES:
            max_mb = MAX_FILE_SIZE_BYTES // (1024 * 1024)
            return False, "", f"Le fichier dépasse la taille maximale autorisée de {max_mb} Mo."

        # 2. Vérification de l'extension
        ext = Path(filename).suffix.lower()
        if ext not in ALLOWED_EXTENSIONS:
            return False, "", f"Format '{ext}' non supporté. Seuls les formats PDF et DOCX sont autorisés."

        # 3. Vérification des magic bytes
        if ext == ".pdf":
            if not file_bytes.startswith(b"%PDF"):
                return False, "", "Fichier PDF corrompu ou invalide (en-tête binaire non conforme)."
            return True, "pdf", ""
        
        elif ext == ".docx":
            # DOCX est une archive ZIP standard (PK\x03\x04 ou PK\x05\x06)
            if not (file_bytes.startswith(b"PK\x03\x04") or file_bytes.startswith(b"PK\x05\x06")):
                return False, "", "Fichier DOCX corrompu ou invalide (signature ZIP absente)."
            return True, "docx", ""

        return False, "", "Format de fichier non pris en charge."

    @staticmethod
    def count_pdf_pages(file_path: Path) -> int:
        """Compte les pages réelles d'un PDF via PyMuPDF côté serveur."""
        try:
            doc = fitz.open(file_path)
            count = len(doc)
            doc.close()
            return count
        except Exception as e:
            raise ValidationError(f"Impossible de lire le document PDF : {str(e)}")

    @staticmethod
    def count_docx_pages(file_path: Path) -> int:
        """
        Estime de manière robuste le nombre de pages d'un fichier Word DOCX.
        Comme le format Word est fluide et ne contient pas de pagination fixe,
        nous combinons :
        - Les sauts de page explicites (pageBreak)
        - Le volume de texte moyen (~350 mots ou 2200 caractères par page standard)
        """
        try:
            doc = Document(file_path)
            total_words = 0
            explicit_breaks = 0

            for p in doc.paragraphs:
                # Vérifier les sauts de page dans les runs
                for run in p.runs:
                    if "w:br" in run._element.xml and 'w:type="page"' in run._element.xml:
                        explicit_breaks += 1
                words = len(p.text.split())
                total_words += words

            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        total_words += len(cell.text.split())

            # Estimation robuste
            estimated_by_words = max(1, (total_words + 300) // 350)
            return max(1, explicit_breaks + 1, estimated_by_words)
        except Exception as e:
            raise ValidationError(f"Impossible de lire le document Word DOCX : {str(e)}")

    @classmethod
    def check_page_limits(cls, page_count: int, is_premium_unlocked: bool = False) -> Tuple[bool, str]:
        """Vérifie si le document respecte le plafond de l'offre gratuite docLove (25 pages)."""
        if page_count > FREE_TIER_MAX_PAGES and not is_premium_unlocked:
            return False, (
                f"Le document contient {page_count} pages, ce qui dépasse la limite "
                f"de l'offre gratuite ({FREE_TIER_MAX_PAGES} pages maximum par document). "
                f"Souscrivez à l'abonnement mensuel docLove à 8 € / mois pour traduire jusqu'à 2 000 pages par mois."
            )
        if page_count > 2000:
            return False, (
                f"Le document contient {page_count} pages, ce qui dépasse le plafond mensuel "
                f"maximal autorisé (2 000 pages). Veuillez scinder le document en plusieurs parties."
            )
        return True, ""
